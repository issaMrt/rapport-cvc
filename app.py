"""Rapports vocaux CVC : dictée -> transcription locale (Whisper) -> Word.
Tout tourne sur ce PC, aucun coût d'API, aucune donnée envoyée à un tiers.
"""
import json
import os
import re
import tempfile
import time
from datetime import datetime
from functools import wraps
from pathlib import Path
from zoneinfo import ZoneInfo

import requests
from docx import Document
from docx.shared import Pt
from flask import (Flask, jsonify, redirect, render_template, request,
                   send_from_directory, session, url_for)
from werkzeug.security import check_password_hash

BASE = Path(__file__).parent
DATA = BASE / "data"
REPORTS = DATA / "rapports"
REPORTS.mkdir(parents=True, exist_ok=True)
CONFIG_FILE = DATA / "config.json"
TZ = ZoneInfo("Europe/Paris")

WHISPER_MODEL = os.environ.get("WHISPER_MODEL", "small")  # tiny/base/small/medium
OLLAMA_URL = os.environ.get("OLLAMA_URL", "http://127.0.0.1:11434")
OLLAMA_MODEL = os.environ.get("OLLAMA_MODEL", "mistral")

if not CONFIG_FILE.exists():
    raise SystemExit("Lancez d'abord : python setup.py (création du mot de passe)")
CONFIG = json.loads(CONFIG_FILE.read_text())

app = Flask(__name__)
app.secret_key = CONFIG["secret_key"]
app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Strict",
    SESSION_COOKIE_SECURE=os.environ.get("INSECURE_LOCAL") != "1",
    MAX_CONTENT_LENGTH=60 * 1024 * 1024,
    PERMANENT_SESSION_LIFETIME=60 * 60 * 12,
)

_whisper = None
_fails = {}  # ip -> (nb échecs, bloqué jusqu'à)


@app.after_request
def headers(resp):
    resp.headers["X-Content-Type-Options"] = "nosniff"
    resp.headers["X-Frame-Options"] = "DENY"
    resp.headers["Referrer-Policy"] = "no-referrer"
    resp.headers["Cache-Control"] = "no-store"
    resp.headers["Content-Security-Policy"] = (
        "default-src 'self'; style-src 'self' 'unsafe-inline'; "
        "script-src 'self' 'unsafe-inline'; media-src 'self' blob:")
    return resp


def login_required(fn):
    @wraps(fn)
    def wrapper(*a, **kw):
        if not session.get("ok"):
            if request.path.startswith("/api/"):
                return jsonify(error="non autorisé"), 401
            return redirect(url_for("login"))
        if request.method == "POST" and request.headers.get("X-Requested-With") != "fetch":
            return jsonify(error="requête refusée"), 403
        return fn(*a, **kw)
    return wrapper


@app.route("/login", methods=["GET", "POST"])
def login():
    ip = request.remote_addr or "?"
    count, until = _fails.get(ip, (0, 0))
    if request.method == "POST":
        if time.time() < until:
            return render_template("login.html", error="Trop d'essais, patientez quelques minutes."), 429
        if check_password_hash(CONFIG["password_hash"], request.form.get("password", "")):
            _fails.pop(ip, None)
            session.clear()
            session.permanent = True
            session["ok"] = True
            return redirect(url_for("index"))
        count += 1
        _fails[ip] = (count, time.time() + 300 if count >= 5 else 0)
        time.sleep(1)
        return render_template("login.html", error="Mot de passe incorrect."), 401
    return render_template("login.html", error=None)


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/")
@login_required
def index():
    return render_template("index.html", technicien=CONFIG.get("technicien", ""))


def get_whisper():
    global _whisper
    if _whisper is None:
        from faster_whisper import WhisperModel
        _whisper = WhisperModel(WHISPER_MODEL, device="cpu", compute_type="int8")
    return _whisper


def charger_audio(path):
    """Décode l'audio en 16 kHz mono, sans passer par faster_whisper.decode_audio
    (incompatible avec certaines versions récentes de PyAV)."""
    import av
    import numpy as np

    resampler = av.AudioResampler(format="s16", layout="mono", rate=16000)
    morceaux = []

    def ajouter(res):
        if res is None:
            return
        for fr in (res if isinstance(res, list) else [res]):
            morceaux.append(fr.to_ndarray().reshape(-1))

    with av.open(path) as conteneur:
        flux = next((s for s in conteneur.streams if s.type == "audio"), None)
        if flux is None:
            raise ValueError("aucune piste audio dans l'enregistrement")
        for frame in conteneur.decode(flux):
            frame.pts = None
            ajouter(resampler.resample(frame))
    ajouter(resampler.resample(None))
    if not morceaux:
        return np.zeros(0, dtype=np.float32)
    return np.concatenate(morceaux).astype(np.float32) / 32768.0


@app.route("/api/transcribe", methods=["POST"])
@login_required
def transcribe():
    f = request.files.get("audio")
    if not f:
        return jsonify(error="audio manquant"), 400
    with tempfile.NamedTemporaryFile(suffix=".bin", delete=False) as tmp:
        f.save(tmp.name)
        path = tmp.name
    try:
        audio = charger_audio(path)
        segments, _ = get_whisper().transcribe(
            audio, language="fr", vad_filter=True, beam_size=5,
            initial_prompt="Rapport d'intervention CVC : chauffage, ventilation, climatisation, "
                           "CTA, VMC, GTB, vanne 3 voies, sonde, groupe froid, débit d'air.")
        text = " ".join(s.text.strip() for s in segments).strip()
    except Exception as e:
        app.logger.exception("Erreur de transcription")
        return jsonify(error=f"transcription impossible ({type(e).__name__}: {e})"), 500
    finally:
        os.unlink(path)
    return jsonify(text=text)


PROMPT = """Tu es rédacteur de rapports techniques pour un technicien d'intégration CVC
(chauffage, ventilation, climatisation, GTB). Le technicien a dicté oralement son compte rendu
d'intervention. Ta mission : le transformer en un rapport professionnel, clair et prêt à être
transmis à un client ou à un responsable.

Règles de rédaction :
- Reformule entièrement : supprime les hésitations, répétitions, tournures orales et digressions
  ("euh", "du coup", "voilà", "genre", etc.).
- Style professionnel et neutre, phrases complètes, voix impersonnelle ou temps du passé
  ("Remplacement de la sonde...", "La vanne a été contrôlée..."). Jamais de "je".
- Utilise le vocabulaire technique exact du métier (CTA, VMC, GTB, régulation, vanne 3 voies,
  sonde de soufflage, servomoteur, automate, etc.).
- Conserve EXACTEMENT les valeurs chiffrées, repères d'équipements, références et noms propres.
- N'invente aucune information, aucune heure, aucune valeur. Si une donnée est absente, laisse vide.
- Classe les activités dans l'ordre chronologique : un titre court (5 à 8 mots) et un détail rédigé.
- "anomalies" : uniquement les dysfonctionnements constatés. "a_faire" : les actions restant à
  réaliser. "materiel" : les pièces ou outils cités.

Réponds UNIQUEMENT avec un JSON valide, de cette forme :
{"resume": "2 à 3 phrases de synthèse",
 "activites": [{"heure": "HH:MM ou vide", "titre": "...", "details": "..."}],
 "materiel": ["..."], "anomalies": ["..."], "a_faire": ["..."]}

Exemple.
Dictée : "alors euh du coup on est arrivé vers huit heures et demie, j'ai regardé la CTA 2 elle soufflait froid, en fait la sonde de soufflage elle était morte donc je l'ai changée, une 10k, après ça marche, 21,4 degrés, par contre la vanne trois voies elle est un peu grippée faudra commander un servomoteur"
Réponse : {"resume": "Intervention sur la CTA 2 pour un défaut de température de soufflage. La sonde défectueuse a été remplacée et le fonctionnement est rétabli. Un défaut de la vanne 3 voies reste à traiter.",
 "activites": [{"heure": "08:30", "titre": "Diagnostic de la CTA 2", "details": "Constat d'un soufflage d'air froid anormal sur la CTA 2."},
 {"heure": "", "titre": "Remplacement de la sonde de soufflage", "details": "La sonde de soufflage défectueuse a été remplacée par une sonde 10k. Après remplacement, la température de soufflage mesurée est de 21,4 °C."}],
 "materiel": ["Sonde de température 10k"], "anomalies": ["Vanne 3 voies partiellement grippée"], "a_faire": ["Commander et remplacer le servomoteur de la vanne 3 voies"]}

Dictée à traiter :
<<TEXTE>>
"""

FILLERS = re.compile(r"\b(euh+|heu+|hum+|ben|bah|bref|voilà|du coup|en fait)\b,?\s*", re.I)


def nettoyer(texte):
    """Nettoyage minimal (repli sans IA) : retire les mots d'hésitation."""
    t = re.sub(r"\s+", " ", FILLERS.sub("", texte)).strip()
    return t[:1].upper() + t[1:] if t else t


def _liste(x):
    return [str(i).strip() for i in x if str(i).strip()] if isinstance(x, list) else []


def structurer(texte):
    """Reformulation professionnelle par l'IA locale (Ollama).
    Repli : nettoyage simple de la dictée si Ollama n'est pas disponible."""
    try:
        r = requests.post(f"{OLLAMA_URL}/api/generate", timeout=300, json={
            "model": OLLAMA_MODEL, "prompt": PROMPT.replace("<<TEXTE>>", texte),
            "format": "json", "stream": False,
            "options": {"temperature": 0.2, "num_ctx": 8192}})
        r.raise_for_status()
        data = json.loads(r.json()["response"])
        activites = []
        brut = data.get("activites")
        for a in (brut if isinstance(brut, list) else []):
            if isinstance(a, dict):
                activites.append({"heure": str(a.get("heure") or ""),
                                  "titre": str(a.get("titre") or ""),
                                  "details": str(a.get("details") or "")})
            elif isinstance(a, str) and a.strip():
                activites.append({"heure": "", "titre": a.strip(), "details": ""})
        if not activites:
            raise ValueError("réponse de l'IA sans activités")
        return {"resume": str(data.get("resume") or ""), "activites": activites,
                "materiel": _liste(data.get("materiel")),
                "anomalies": _liste(data.get("anomalies")),
                "a_faire": _liste(data.get("a_faire"))}, "ia"
    except Exception as e:
        app.logger.warning("IA locale indisponible, repli simple : %s", e)
        phrases = [p.strip() for p in re.split(r"(?<=[.!?])\s+", nettoyer(texte)) if p.strip()]
        return {"resume": "", "materiel": [], "anomalies": [], "a_faire": [],
                "activites": [{"heure": "", "titre": p, "details": ""} for p in phrases]}, "simple"


def slug(s):
    import unicodedata
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")[:40] or "rapport"


def faire_docx(meta, rapport, texte):
    now = datetime.now(TZ)
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading("Rapport d'intervention CVC", 0)

    t = doc.add_table(rows=0, cols=2)
    t.style = "Light Grid Accent 1"
    for k, v in [("Date", meta.get("date") or now.strftime("%d/%m/%Y")),
                 ("Heure de début", meta.get("debut", "")),
                 ("Heure de fin", meta.get("fin", "")),
                 ("Chantier / site", meta.get("chantier", "")),
                 ("Technicien", meta.get("technicien", "")),
                 ("Rédigé le", now.strftime("%d/%m/%Y à %H:%M"))]:
        row = t.add_row().cells
        row[0].text, row[1].text = k, v

    if rapport.get("resume"):
        doc.add_heading("Résumé", 1)
        doc.add_paragraph(rapport["resume"])

    doc.add_heading("Détail des activités", 1)
    if not rapport["activites"]:
        doc.add_paragraph("Aucune activité renseignée.")
    for a in rapport["activites"]:
        p = doc.add_paragraph(style="List Number")
        heure = f"[{a['heure']}] " if a.get("heure") else ""
        r = p.add_run(heure + (a.get("titre") or ""))
        r.bold = True
        if a.get("details"):
            p.add_run(" — " + a["details"])

    for titre, cle in [("Matériel / pièces", "materiel"), ("Anomalies constatées", "anomalies"),
                       ("Reste à faire", "a_faire")]:
        if rapport.get(cle):
            doc.add_heading(titre, 1)
            for x in rapport[cle]:
                doc.add_paragraph(str(x), style="List Bullet")

    doc.add_page_break()
    doc.add_heading("Annexe : transcription brute", 1)
    doc.add_paragraph(texte)

    name = f"{now:%Y-%m-%d_%Hh%M}_{slug(meta.get('chantier', ''))}.docx"
    doc.save(REPORTS / name)
    return name


@app.route("/api/rapport", methods=["POST"])
@login_required
def rapport():
    d = request.get_json(force=True)
    texte = (d.get("texte") or "").strip()
    if not texte:
        return jsonify(error="transcription vide"), 400
    structure, mode = structurer(texte)
    name = faire_docx(d.get("meta", {}), structure, texte)
    return jsonify(fichier=name, mode=mode, rapport=structure)


@app.route("/api/rapports")
@login_required
def liste():
    files = sorted(REPORTS.glob("*.docx"), reverse=True)
    return jsonify([f.name for f in files][:50])


@app.route("/rapports/<path:name>")
@login_required
def telecharger(name):
    return send_from_directory(REPORTS, name, as_attachment=True)


if __name__ == "__main__":
    # Écoute en local uniquement : l'accès distant passe par Tailscale (voir README)
    from waitress import serve
    serve(app, host="127.0.0.1", port=8000, threads=4)
